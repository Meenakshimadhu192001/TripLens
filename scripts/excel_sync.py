from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
WORKBOOK_NAMES = (
    "Package_Details_India(final).xlsx",
    "PPackage_Details__.xlsx",
)


def _headers(sheet):
    return {cell.value: index for index, cell in enumerate(sheet[1], start=1) if cell.value}


def _find_or_append_row(sheet, package_id):
    columns = _headers(sheet)
    package_column = columns.get("package_id")
    if not package_column:
        raise ValueError(f"Sheet {sheet.title} has no package_id column")

    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row, package_column).value or "").strip() == package_id:
            return row

    return sheet.max_row + 1


def _write_row(sheet, values, package_id):
    columns = _headers(sheet)
    row = _find_or_append_row(sheet, package_id)
    for key, value in values.items():
        column = columns.get(key)
        if column:
            sheet.cell(row, column, value)


def _replace_child_rows(sheet, package_id, rows):
    columns = _headers(sheet)
    package_column = columns.get("package_id")
    if not package_column:
        raise ValueError(f"Sheet {sheet.title} has no package_id column")

    for row in range(sheet.max_row, 1, -1):
        if str(sheet.cell(row, package_column).value or "").strip() == package_id:
            sheet.delete_rows(row, 1)

    for values in rows:
        row = sheet.max_row + 1
        for key, value in values.items():
            column = columns.get(key)
            if column:
                sheet.cell(row, column, value)


def _ensure_sheet(workbook, title, headers):
    if title in workbook.sheetnames:
        sheet = workbook[title]
        if sheet.max_row == 1 and all(cell.value is None for cell in sheet[1]):
            for index, value in enumerate(headers, start=1):
                sheet.cell(1, index, value)
        return sheet

    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    return sheet


def _sync_workbook(path, payload):
    workbook = load_workbook(path)
    package_id = payload.package_id.strip().upper()

    _write_row(
        workbook["Packages"],
        {
            "package_id": package_id,
            "agency_name": payload.agency_name.strip(),
            "package_name": payload.package_name.strip(),
            "source_url_or_doc": "portal",
            "data_source": "packager_submission",
            "destinations": payload.destinations.strip(),
            "start_location": payload.start_location.strip(),
            "duration_days": payload.duration_days,
            "duration_nights": payload.duration_nights,
            "price": payload.price,
            "price_detail": str(payload.price),
            "tier_range_exists": "FALSE",
            "transport_type": payload.transport_type.strip(),
            "theme": payload.theme.strip(),
            "suited_for": "",
            "itinerary_pace": payload.itinerary_pace.strip(),
            "customizable": "TRUE",
            "agency_contact": payload.agency_contact.strip(),
        },
        package_id,
    )

    _replace_child_rows(
        workbook["Itinerary_Days"],
        package_id,
        [
            {
                "package_id": package_id,
                "day_number": day.day_number,
                "stops": day.stops.strip(),
                "activities": day.activities.strip(),
                "activity_type": day.activity_type.strip(),
                "meals_included_today": day.meals_included_today.strip(),
                "flagged_for_verification": "FALSE",
                "notes": "",
            }
            for day in payload.itinerary
        ],
    )

    _replace_child_rows(
        workbook["Accommodation"],
        package_id,
        [
            {
                "package_id": package_id,
                "destination": stay.destination.strip(),
                "accommodation_category": stay.accommodation_category.strip(),
                "hotel_name": stay.hotel_name.strip(),
                "hotel_guaranteed": stay.hotel_guaranteed,
            }
            for stay in payload.accommodation
        ],
    )

    inclusions = _ensure_sheet(workbook, "Inclusions", ["package_id", "inclusion_item"])
    exclusions = _ensure_sheet(workbook, "Exclusions", ["package_id", "exclusion_item"])
    _replace_child_rows(
        inclusions,
        package_id,
        [{"package_id": package_id, "inclusion_item": item.strip()} for item in payload.inclusions if item.strip()],
    )
    _replace_child_rows(
        exclusions,
        package_id,
        [{"package_id": package_id, "exclusion_item": item.strip()} for item in payload.exclusions if item.strip()],
    )

    workbook.save(path)


def sync_package_to_excel(payload):
    paths = [PROJECT_ROOT / name for name in WORKBOOK_NAMES if (PROJECT_ROOT / name).exists()]
    if not paths:
        raise FileNotFoundError("No package-details workbook was found")

    for path in paths:
        _sync_workbook(path, payload)

    return [path.name for path in paths]
